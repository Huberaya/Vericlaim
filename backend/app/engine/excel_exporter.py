"""Générateur de Reporting Excel Décisionnel Multi-feuilles (VeriClaim AI).

Ce module compile les résultats d'audit réglementaire et d'évaluation de conformité
en un classeur Excel OOXML (.xlsx) multi-onglets hautement structuré et auditable,
conçu pour les directions juridiques, RSE, achats et comités de direction.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.schemas import CatalogBatchResponse, EvaluationResponse

# Palette chromatique VeriClaim AI
COLOR_NAVY = "0F172A"       # En-tête principal & Brand
COLOR_NAVY_LIGHT = "1E293B" # Sous-sections
COLOR_BLUE = "1E40AF"       # Références juridiques & Liens
COLOR_SLATE = "475569"      # Métadonnées & Textes secondaires
COLOR_LIGHT_BG = "F8FAFC"   # Fonds alternés
COLOR_BORDER = "CBD5E1"     # Bordures de cellules standard
COLOR_BORDER_THICK = "94A3B8"

# Couleurs de statut juridique
COLOR_RED_FILL = "FEE2E2"
COLOR_RED_TEXT = "991B1B"
COLOR_GREEN_FILL = "DCFCE7"
COLOR_GREEN_TEXT = "166534"
COLOR_AMBER_FILL = "FEF3C7"
COLOR_AMBER_TEXT = "92400E"
COLOR_GRAY_FILL = "F1F5F9"
COLOR_GRAY_TEXT = "475569"


def _create_border(
    top: str = COLOR_BORDER,
    bottom: str = COLOR_BORDER,
    left: str = COLOR_BORDER,
    right: str = COLOR_BORDER,
    style: str = "thin",
) -> Border:
    return Border(
        top=Side(style=style, color=top),
        bottom=Side(style=style, color=bottom),
        left=Side(style=style, color=left),
        right=Side(style=style, color=right),
    )


def _get_status_colors(status: str) -> tuple[str, str, str]:
    """Retourne (libellé_fr, fill_color, text_color)."""
    s = str(status).upper()
    if "NON_COMPLIANT" in s:
        return "NON CONFORME — INFRACTION RETENUE", COLOR_RED_FILL, COLOR_RED_TEXT
    if "COMPLIANT" in s:
        return "CONFORME AU CADRE CONTRÔLÉ", COLOR_GREEN_FILL, COLOR_GREEN_TEXT
    if "CONDITIONAL_REJECT" in s:
        return "PUBLICATION À SUSPENDRE — PREUVES REQUISES", COLOR_AMBER_FILL, COLOR_AMBER_TEXT
    if "REVIEW_REQUIRED" in s:
        return "REVUE JURIDIQUE / PROBATOIRE REQUISE", COLOR_AMBER_FILL, COLOR_AMBER_TEXT
    if "NO_CLAIMS_DETECTED" in s:
        return "AUCUNE ALLÉGATION DÉTECTÉE", COLOR_GRAY_FILL, COLOR_GRAY_TEXT
    return s.replace("_", " "), COLOR_GRAY_FILL, COLOR_GRAY_TEXT


def _auto_fit_columns(ws, min_width: int = 12, max_width: int = 60) -> None:
    """Ajuste la largeur des colonnes intelligemment."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Ne pas considérer les cellules fusionnées sur plusieurs colonnes
            if cell.coordinate in ws.merged_cells:
                continue
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 3, max_width))


def generate_audit_excel(evaluation: EvaluationResponse) -> bytes:
    """Génère le classeur Excel multi-feuilles d'audit décisionnel VeriClaim AI."""
    wb = openpyxl.Workbook()
    # Supprimer la feuille par défaut après création
    default_sheet = wb.active

    thin_border = _create_border()
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    navy_fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
    sub_fill = PatternFill(start_color=COLOR_NAVY_LIGHT, end_color=COLOR_NAVY_LIGHT, fill_type="solid")
    alt_fill = PatternFill(start_color=COLOR_LIGHT_BG, end_color=COLOR_LIGHT_BG, fill_type="solid")

    status_label, status_fill, status_text = _get_status_colors(evaluation.overall_compliance.value)

    # -------------------------------------------------------------------------
    # FEUILLE 1 : SYNTHÈSE DIRECTION
    # -------------------------------------------------------------------------
    ws_summary = wb.create_sheet(title="Synthèse Direction")
    ws_summary.sheet_properties.tabColor = COLOR_NAVY
    ws_summary.views.sheetView[0].showGridLines = True

    # Titre Header Corporate
    ws_summary.merge_cells("A1:G2")
    title_cell = ws_summary["A1"]
    title_cell.value = "VERICLAIM AI · RAPPORT DÉCISIONNEL D'AUDIT RÉGLEMENTAIRE"
    title_cell.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    title_cell.fill = navy_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws_summary["A3"].value = (
        f"Généré le {datetime.now(timezone.utc).strftime('%d/%m/%Y à %H:%M UTC')} — "
        f"Audit ID: {evaluation.audit_trail.audit_id} — Empreinte SHA-256 certifiée"
    )
    ws_summary["A3"].font = Font(name="Calibri", size=9, italic=True, color=COLOR_SLATE)
    ws_summary.merge_cells("A3:G3")

    # Cadre Métadonnées de l'audit
    meta_rows = [
        ("Identifiant d'Audit", evaluation.audit_trail.audit_id),
        ("Date d'effet / Évaluation", str(evaluation.audit_trail.as_of_date)),
        ("Juridiction d'évaluation", "France / Union Européenne (AGEC · EmpCo 2024/825 · C. conso)"),
        ("Moteur & RuleBook", f"{evaluation.audit_trail.engine_version} / {evaluation.audit_trail.rulebook_version}"),
        ("Organisation / Tenant", evaluation.audit_trail.tenant_id),
        ("Méthode d'extraction", evaluation.audit_trail.extraction_method),
    ]
    for idx, (lbl, val) in enumerate(meta_rows, start=5):
        ws_summary[f"A{idx}"] = lbl
        ws_summary[f"A{idx}"].font = Font(name="Calibri", size=10, bold=True, color=COLOR_NAVY)
        ws_summary[f"A{idx}"].fill = alt_fill
        ws_summary[f"A{idx}"].border = thin_border

        ws_summary[f"B{idx}"] = val
        ws_summary[f"B{idx}"].font = Font(name="Calibri", size=10)
        ws_summary[f"B{idx}"].border = thin_border
        ws_summary.merge_cells(f"B{idx}:G{idx}")

    # Section KPIs Décisionnels (Lignes 12 à 15)
    ws_summary["A12"].value = "INDICATEURS CLÉS DE GOUVERNANCE & CONFORMITÉ"
    ws_summary["A12"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws_summary["A12"].fill = sub_fill
    ws_summary["A12"].alignment = Alignment(vertical="center", indent=1)
    ws_summary.merge_cells("A12:G12")

    kpi_cards = [
        ("STATUT DE CONFORMITÉ", status_label, status_fill, status_text),
        ("SCORE DE RISQUE", f"{evaluation.risk_score} / 100", COLOR_RED_FILL if evaluation.risk_score > 50 else COLOR_AMBER_FILL if evaluation.risk_score > 20 else COLOR_GREEN_FILL, COLOR_RED_TEXT if evaluation.risk_score > 50 else COLOR_AMBER_TEXT if evaluation.risk_score > 20 else COLOR_GREEN_TEXT),
        ("ALLÉGATIONS DÉTECTÉES", evaluation.detected_claims_count, COLOR_GRAY_FILL, COLOR_NAVY),
        ("INFRACTIONS RETENUES", evaluation.violations_count, COLOR_RED_FILL if evaluation.violations_count > 0 else COLOR_GREEN_FILL, COLOR_RED_TEXT if evaluation.violations_count > 0 else COLOR_GREEN_TEXT),
        ("RÉSERVES PROBATOIRES", evaluation.conditional_findings_count, COLOR_AMBER_FILL if evaluation.conditional_findings_count > 0 else COLOR_GRAY_FILL, COLOR_AMBER_TEXT if evaluation.conditional_findings_count > 0 else COLOR_NAVY),
    ]

    col_letters = [("A", "B"), ("C", "C"), ("D", "D"), ("E", "E"), ("F", "G")]
    for (label, val, f_col, t_col), (c_start, c_end) in zip(kpi_cards, col_letters):
        ws_summary[f"{c_start}13"] = label
        ws_summary[f"{c_start}13"].font = Font(name="Calibri", size=8.5, bold=True, color=COLOR_SLATE)
        ws_summary[f"{c_start}13"].alignment = Alignment(horizontal="center", vertical="center")
        ws_summary[f"{c_start}13"].border = thin_border
        if c_start != c_end:
            ws_summary.merge_cells(f"{c_start}13:{c_end}13")

        ws_summary[f"{c_start}14"] = val
        ws_summary[f"{c_start}14"].font = Font(name="Calibri", size=13, bold=True, color=t_col)
        ws_summary[f"{c_start}14"].fill = PatternFill(start_color=f_col, end_color=f_col, fill_type="solid")
        ws_summary[f"{c_start}14"].alignment = Alignment(horizontal="center", vertical="center")
        ws_summary[f"{c_start}14"].border = thin_border
        if c_start != c_end:
            ws_summary.merge_cells(f"{c_start}14:{c_end}14")

    # Section Exposition Financière & Sanctions (Lignes 16 à 24)
    ws_summary["A16"].value = "EXPOSITION JURIDIQUE & CHIFSTAGE DU RISQUE FINANCIER"
    ws_summary["A16"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws_summary["A16"].fill = sub_fill
    ws_summary["A16"].alignment = Alignment(vertical="center", indent=1)
    ws_summary.merge_cells("A16:G16")

    headers_exposure = ["Catégorie", "Fondement Légal", "Plafond Personne Morale", "Plafond Personne Physique", "Risque Proportionnel", "Qualification / Risque"]
    for idx_col, h in enumerate(headers_exposure, start=1):
        c = ws_summary.cell(row=17, column=idx_col, value=h)
        c.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color=COLOR_SLATE, end_color=COLOR_SLATE, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    curr_row = 18
    if evaluation.exposure_matrix and evaluation.exposure_matrix.items:
        for item in evaluation.exposure_matrix.items:
            ws_summary.cell(row=curr_row, column=1, value=item.category.value).border = thin_border
            ws_summary.cell(row=curr_row, column=2, value=item.legal_basis or item.title).border = thin_border
            
            c_moral = ws_summary.cell(row=curr_row, column=3, value=float(item.max_legal_person_eur) if item.max_legal_person_eur else "Sur CA / Dépenses")
            c_moral.number_format = "#,##0 €" if item.max_legal_person_eur else "@"
            c_moral.border = thin_border

            c_phys = ws_summary.cell(row=curr_row, column=4, value=float(item.max_natural_person_eur) if item.max_natural_person_eur else "-")
            c_phys.number_format = "#,##0 €" if item.max_natural_person_eur else "@"
            c_phys.border = thin_border

            ws_summary.cell(row=curr_row, column=5, value="Oui (jusqu'à 80% dépenses de pub)" if item.may_scale_to_advertising_spend else "Non").border = thin_border
            ws_summary.cell(row=curr_row, column=6, value=item.note).border = thin_border
            curr_row += 1
    else:
        ws_summary.cell(row=curr_row, column=1, value="Aucun risque financier direct caractérisé.").border = thin_border
        ws_summary.merge_cells(f"A{curr_row}:F{curr_row}")
        curr_row += 1

    # Avis de la Direction de l'Audit
    curr_row += 1
    ws_summary.cell(row=curr_row, column=1, value="AVIS STRATÉGIQUE & RECOMMANDATION IMMÉDIATE").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws_summary.cell(row=curr_row, column=1).fill = sub_fill
    ws_summary.merge_cells(f"A{curr_row}:G{curr_row}")
    curr_row += 1

    recommendation_text = (
        "🔴 SUSPENSION IMMÉDIATE DU SUPPORT : Des mentions strictement interdites par la loi (ex. AGEC L. 541-9-1) "
        "ou trompeuses ont été constatées. Retirer le produit du circuit ou corriger le packaging avant toute commercialisation."
        if evaluation.overall_compliance.value == "NON_COMPLIANT"
        else "🟡 PUBLICATION CONDITIONNÉE : Les allégations requièrent la consolidation préalable des pièces justificatives "
        "(rapport ACV conforme ISO 14040/44 ou enregistrement officiel d'écolabel Type I) avant diffusion publique."
        if evaluation.overall_compliance.value in ("CONDITIONAL_REJECT", "REVIEW_REQUIRED")
        else "🟢 VALIDATION CONFORME : Aucune non-conformité relevée dans le périmètre des règles auditées. "
        "Conserver l'attestation cryptographique et le registre dans le dossier d'opposabilité de l'entreprise."
    )
    rec_cell = ws_summary.cell(row=curr_row, column=1, value=recommendation_text)
    rec_cell.font = Font(name="Calibri", size=10, bold=True, color=status_text)
    rec_cell.fill = PatternFill(start_color=status_fill, end_color=status_fill, fill_type="solid")
    rec_cell.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    rec_cell.border = thin_border
    ws_summary.merge_cells(f"A{curr_row}:G{curr_row + 1}")
    ws_summary.row_dimensions[curr_row].height = 28
    ws_summary.row_dimensions[curr_row + 1].height = 28

    _auto_fit_columns(ws_summary, min_width=14, max_width=45)

    # -------------------------------------------------------------------------
    # FEUILLE 2 : REGISTRE DES ALLÉGATIONS AUDITÉES
    # -------------------------------------------------------------------------
    ws_claims = wb.create_sheet(title="Registre Allégations")
    ws_claims.sheet_properties.tabColor = COLOR_BLUE
    ws_claims.views.sheetView[0].showGridLines = True
    ws_claims.freeze_panes = "A3"

    ws_claims.merge_cells("A1:K1")
    ws_claims["A1"] = "REGISTRE JURIDIQUE DES ALLÉGATIONS DÉTECTÉES & ÉVALUATIONS DÉTERMINISTES"
    ws_claims["A1"].font = header_font
    ws_claims["A1"].fill = navy_fill
    ws_claims["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_claims.row_dimensions[1].height = 30

    claim_headers = [
        "Réf / ID",
        "Allégation Détectée (Extrait)",
        "Typologie Réglementaire",
        "Règle RuleBook",
        "Base Légale (Article de Loi)",
        "Force Juridique",
        "Verdict Moteur",
        "Infraction Légale ?",
        "Safe Harbor ?",
        "Sévérité Risque",
        "Preuves Exigées",
    ]
    ws_claims.row_dimensions[2].height = 26
    for idx_col, h in enumerate(claim_headers, start=1):
        c = ws_claims.cell(row=2, column=idx_col, value=h)
        c.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
        c.fill = sub_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

    r_idx = 3
    for eval_item in evaluation.evaluations:
        ws_claims.row_dimensions[r_idx].height = 32
        _, f_col, t_col = _get_status_colors(eval_item.verdict.value)

        ws_claims.cell(row=r_idx, column=1, value=eval_item.claim_id).border = thin_border
        
        c_text = ws_claims.cell(row=r_idx, column=2, value=f"« {eval_item.claim_text} »")
        c_text.font = Font(name="Calibri", size=9.5, italic=True)
        c_text.alignment = Alignment(wrap_text=True, vertical="center")
        c_text.border = thin_border

        ws_claims.cell(row=r_idx, column=3, value=eval_item.claim_type.value).border = thin_border
        ws_claims.cell(row=r_idx, column=4, value=eval_item.rule_title).border = thin_border

        c_law = ws_claims.cell(row=r_idx, column=5, value=eval_item.law_reference)
        c_law.font = Font(name="Calibri", size=9, bold=True, color=COLOR_BLUE)
        c_law.border = thin_border

        ws_claims.cell(row=r_idx, column=6, value=eval_item.legal_force.value).border = thin_border

        c_verd = ws_claims.cell(row=r_idx, column=7, value=eval_item.verdict.value)
        c_verd.font = Font(name="Calibri", size=9, bold=True, color=t_col)
        c_verd.fill = PatternFill(start_color=f_col, end_color=f_col, fill_type="solid")
        c_verd.alignment = Alignment(horizontal="center", vertical="center")
        c_verd.border = thin_border

        c_inf = ws_claims.cell(row=r_idx, column=8, value="OUI — INFRACTION" if eval_item.is_legal_violation else "NON")
        c_inf.font = Font(name="Calibri", size=9, bold=True, color=COLOR_RED_TEXT if eval_item.is_legal_violation else COLOR_GREEN_TEXT)
        c_inf.alignment = Alignment(horizontal="center", vertical="center")
        c_inf.border = thin_border

        ws_claims.cell(row=r_idx, column=9, value="ACCORDÉ" if eval_item.safe_harbor_applicable else "NON").border = thin_border
        
        c_sev = ws_claims.cell(row=r_idx, column=10, value=eval_item.severity.value)
        c_sev.alignment = Alignment(horizontal="center", vertical="center")
        c_sev.border = thin_border

        c_req = ws_claims.cell(row=r_idx, column=11, value=" ; ".join(eval_item.required_evidence) if eval_item.required_evidence else "Aucune preuve requise")
        c_req.alignment = Alignment(wrap_text=True, vertical="center")
        c_req.border = thin_border

        r_idx += 1

    if len(evaluation.evaluations) == 0:
        c_none = ws_claims.cell(row=r_idx, column=1, value="Aucune allégation environnementale identifiée par le lexique réglementaire.")
        c_none.border = thin_border
        ws_claims.merge_cells(f"A{r_idx}:K{r_idx}")

    _auto_fit_columns(ws_claims, min_width=12, max_width=50)

    # -------------------------------------------------------------------------
    # FEUILLE 3 : PLAN DE REMÉDIATION & CLAUSES CONTRACTUELLES
    # -------------------------------------------------------------------------
    ws_remed = wb.create_sheet(title="Plan Remédiation")
    ws_remed.sheet_properties.tabColor = "059669" # Emerald green
    ws_remed.views.sheetView[0].showGridLines = True
    ws_remed.freeze_panes = "A3"

    ws_remed.merge_cells("A1:F1")
    ws_remed["A1"] = "PLAN D'ACTION OPÉRATIONNEL & CLAUSES DE REMÉDIATION CONTRACTUELLE"
    ws_remed["A1"].font = header_font
    ws_remed["A1"].fill = navy_fill
    ws_remed["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_remed.row_dimensions[1].height = 30

    remed_headers = [
        "Réf Allégation",
        "Texte Original à Corriger",
        "Explication Juridique Acheteur / Marketing",
        "Formulation de Remplacement Recommandée",
        "Clause Contractuelle de Substitution (Avenant)",
        "Actions Immédiates Obligatoires",
    ]
    ws_remed.row_dimensions[2].height = 26
    for idx_col, h in enumerate(remed_headers, start=1):
        c = ws_remed.cell(row=2, column=idx_col, value=h)
        c.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
        c.fill = sub_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

    rem_row = 3
    for eval_item in evaluation.evaluations:
        ws_remed.row_dimensions[rem_row].height = 55
        
        ws_remed.cell(row=rem_row, column=1, value=eval_item.claim_id).border = thin_border
        
        c_orig = ws_remed.cell(row=rem_row, column=2, value=f"« {eval_item.claim_text} »")
        c_orig.font = Font(name="Calibri", size=9.5, strike=eval_item.is_legal_violation, color=COLOR_RED_TEXT if eval_item.is_legal_violation else COLOR_NAVY)
        c_orig.alignment = Alignment(wrap_text=True, vertical="top")
        c_orig.border = thin_border

        c_exp = ws_remed.cell(row=rem_row, column=3, value=eval_item.remediation.buyer_explanation)
        c_exp.alignment = Alignment(wrap_text=True, vertical="top")
        c_exp.border = thin_border

        c_rew = ws_remed.cell(row=rem_row, column=4, value=eval_item.remediation.recommended_rewrite)
        c_rew.font = Font(name="Calibri", size=9.5, bold=True, color=COLOR_GREEN_TEXT)
        c_rew.alignment = Alignment(wrap_text=True, vertical="top")
        c_rew.border = thin_border

        c_clause = ws_remed.cell(row=rem_row, column=5, value=eval_item.remediation.supplier_contract_clause)
        c_clause.font = Font(name="Calibri", size=8.5, italic=True)
        c_clause.alignment = Alignment(wrap_text=True, vertical="top")
        c_clause.border = thin_border

        actions_str = " ; \n".join(eval_item.remediation.required_actions) if eval_item.remediation.required_actions else "Aucune action bloquante"
        c_act = ws_remed.cell(row=rem_row, column=6, value=actions_str)
        c_act.alignment = Alignment(wrap_text=True, vertical="top")
        c_act.border = thin_border

        rem_row += 1

    if len(evaluation.evaluations) == 0:
        c_none = ws_remed.cell(row=rem_row, column=1, value="Aucune remédiation nécessaire.")
        c_none.border = thin_border
        ws_remed.merge_cells(f"A{rem_row}:F{rem_row}")

    _auto_fit_columns(ws_remed, min_width=15, max_width=45)

    # -------------------------------------------------------------------------
    # FEUILLE 4 : MATRICE DE PREUVES & TRAÇABILITÉ (LEDGER)
    # -------------------------------------------------------------------------
    ws_ledger = wb.create_sheet(title="Matrice Traçabilité")
    ws_ledger.sheet_properties.tabColor = "7C3AED" # Purple
    ws_ledger.views.sheetView[0].showGridLines = True

    ws_ledger.merge_cells("A1:D1")
    ws_ledger["A1"] = "REGISTRE CRYPTOGRAPHIQUE & SCEAU D'OPPOSABILITÉ JURIDIQUE (LEDGER)"
    ws_ledger["A1"].font = header_font
    ws_ledger["A1"].fill = navy_fill
    ws_ledger["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_ledger.row_dimensions[1].height = 30

    ws_ledger["A2"] = (
        "Ce registre atteste de l'intégrité formelle de l'évaluation et garantit l'inviolabilité "
        "des calculs face aux contrôles de la DGCCRF et des auditeurs indépendants."
    )
    ws_ledger["A2"].font = Font(name="Calibri", size=9, italic=True, color=COLOR_SLATE)
    ws_ledger.merge_cells("A2:D2")

    trail = evaluation.audit_trail
    ledger_records = [
        ("Paramètre Cryptographique", "Valeur Enregistrée", "Algorithme", "Statut Opposabilité"),
        ("Identifiant Universel de l'Audit (UUID)", trail.audit_id, "UUIDv4", "CERTIFIÉ"),
        ("Organisation / Tenant ID", trail.tenant_id, "Multi-Tenant Isolation", "CERTIFIÉ"),
        ("Date d'horodatage UTC", trail.evaluated_at_utc.isoformat(), "ISO-8601 UTC", "CERTIFIÉ"),
        ("Date légale d'opposabilité", str(trail.as_of_date), "Calendrier Civil", "CERTIFIÉ"),
        ("Empreinte du Texte Source", trail.source_sha256, "SHA-256 (FIPS 180-4)", "IMMUABLE"),
        ("Empreinte du Document Original", trail.document_sha256 or "Non fourni (saisie directe)", "SHA-256 (FIPS 180-4)", "CERTIFIÉ" if trail.document_sha256 else "N/A"),
        ("Empreinte du Manifeste de Preuves", trail.evidence_manifest_sha256, "SHA-256", "VÉRIFIÉ"),
        ("Empreinte d'Intégrité du Rapport", trail.report_sha256, "SHA-256 (FIPS 180-4)", "IMMUABLE"),
        ("Bloc Précédent (Chain Link)", trail.previous_record_hash or "0" * 64, "SHA-256 Chain", "VÉRIFIÉ"),
        ("Empreinte de Bloc du Ledger (Record Hash)", trail.record_hash, "SHA-256 Immutable Block", "SCELLÉ & OPPOSABLE"),
    ]

    for row_num, record in enumerate(ledger_records, start=4):
        ws_ledger.row_dimensions[row_num].height = 24
        is_header = row_num == 4
        for col_num, val in enumerate(record, start=1):
            cell = ws_ledger.cell(row=row_num, column=col_num, value=val)
            if is_header:
                cell.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
                cell.fill = sub_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.border = thin_border
                if col_num in (2,):
                    cell.font = Font(name="Consolas", size=9)
                elif col_num == 4:
                    cell.font = Font(name="Calibri", size=9, bold=True, color=COLOR_GREEN_TEXT)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = Font(name="Calibri", size=9.5)

    _auto_fit_columns(ws_ledger, min_width=18, max_width=68)

    # Supprimer la feuille par défaut si présente
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    # Activer la première feuille à l'ouverture
    wb.active = ws_summary

    output_stream = io.BytesIO()
    wb.save(output_stream)
    return output_stream.getvalue()


def generate_catalog_batch_excel(batch: CatalogBatchResponse) -> bytes:
    """Génère un classeur Excel pour un audit de catalogue complet (batch 500+ SKUs)."""
    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Synthèse Catalogue"
    ws_summary.sheet_properties.tabColor = COLOR_NAVY
    ws_summary.views.sheetView[0].showGridLines = True

    thin_border = _create_border()
    navy_fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
    sub_fill = PatternFill(start_color=COLOR_NAVY_LIGHT, end_color=COLOR_NAVY_LIGHT, fill_type="solid")

    ws_summary.merge_cells("A1:G2")
    ws_summary["A1"] = "VERICLAIM AI · AUDIT RÉGLEMENTAIRE DE CATALOGUE PRODUITS"
    ws_summary["A1"].font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    ws_summary["A1"].fill = navy_fill
    ws_summary["A1"].alignment = Alignment(horizontal="center", vertical="center")

    # KPIs Catalogue
    kpi_items = [
        ("TOTAL PRODUITS (SKUs)", batch.total_items, COLOR_GRAY_FILL, COLOR_NAVY),
        ("PRODUITS CONFORMES", batch.compliant_items, COLOR_GREEN_FILL, COLOR_GREEN_TEXT),
        ("PRODUITS NON CONFORMES", batch.non_compliant_items, COLOR_RED_FILL, COLOR_RED_TEXT),
        ("TAUX DE CONFORMITÉ", f"{batch.compliance_rate_pct:.1f} %", COLOR_AMBER_FILL if batch.compliance_rate_pct < 80 else COLOR_GREEN_FILL, COLOR_AMBER_TEXT if batch.compliance_rate_pct < 80 else COLOR_GREEN_TEXT),
        ("SCORE RISQUE MOYEN", f"{batch.average_risk_score:.1f} / 100", COLOR_RED_FILL if batch.average_risk_score > 40 else COLOR_GREEN_FILL, COLOR_RED_TEXT if batch.average_risk_score > 40 else COLOR_GREEN_TEXT),
        ("PLAFOND AMENDES THÉORIQUE", f"{batch.total_fines_ceiling_eur:,.0f} €".replace(",", " "), COLOR_RED_FILL, COLOR_RED_TEXT),
    ]

    col_mappings = [("A", "A"), ("B", "B"), ("C", "C"), ("D", "D"), ("E", "E"), ("F", "G")]
    for (lbl, val, f_col, t_col), (c_s, c_e) in zip(kpi_items, col_mappings):
        ws_summary[f"{c_s}4"] = lbl
        ws_summary[f"{c_s}4"].font = Font(name="Calibri", size=8.5, bold=True, color=COLOR_SLATE)
        ws_summary[f"{c_s}4"].alignment = Alignment(horizontal="center", vertical="center")
        ws_summary[f"{c_s}4"].border = thin_border
        if c_s != c_e:
            ws_summary.merge_cells(f"{c_s}4:{c_e}4")

        ws_summary[f"{c_s}5"] = val
        ws_summary[f"{c_s}5"].font = Font(name="Calibri", size=13, bold=True, color=t_col)
        ws_summary[f"{c_s}5"].fill = PatternFill(start_color=f_col, end_color=f_col, fill_type="solid")
        ws_summary[f"{c_s}5"].alignment = Alignment(horizontal="center", vertical="center")
        ws_summary[f"{c_s}5"].border = thin_border
        if c_s != c_e:
            ws_summary.merge_cells(f"{c_s}5:{c_e}5")

    # Tableau des SKUs
    ws_skus = wb.create_sheet(title="Détail par SKU")
    ws_skus.sheet_properties.tabColor = COLOR_BLUE
    ws_skus.views.sheetView[0].showGridLines = True
    ws_skus.freeze_panes = "A2"

    sku_headers = [
        "SKU",
        "Titre Produit",
        "Fournisseur",
        "Statut Juridique",
        "Score Risque",
        "Allégations Détectées",
        "Infractions Constatées",
        "Plafond d'Amendes (€)",
        "Synthèse d'Audit",
    ]
    for idx_col, h in enumerate(sku_headers, start=1):
        c = ws_skus.cell(row=1, column=idx_col, value=h)
        c.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
        c.fill = sub_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    for row_idx, res in enumerate(batch.results, start=2):
        _, f_col, t_col = _get_status_colors(res.overall_compliance.value)

        ws_skus.cell(row=row_idx, column=1, value=res.sku).border = thin_border
        ws_skus.cell(row=row_idx, column=2, value=res.title).border = thin_border
        ws_skus.cell(row=row_idx, column=3, value=res.supplier_name or "-").border = thin_border

        c_stat = ws_skus.cell(row=row_idx, column=4, value=res.overall_compliance.value)
        c_stat.font = Font(name="Calibri", size=9, bold=True, color=t_col)
        c_stat.fill = PatternFill(start_color=f_col, end_color=f_col, fill_type="solid")
        c_stat.alignment = Alignment(horizontal="center", vertical="center")
        c_stat.border = thin_border

        c_risk = ws_skus.cell(row=row_idx, column=5, value=res.risk_score)
        c_risk.alignment = Alignment(horizontal="center", vertical="center")
        c_risk.border = thin_border

        ws_skus.cell(row=row_idx, column=6, value=res.detected_claims_count).border = thin_border
        ws_skus.cell(row=row_idx, column=7, value=res.violations_count).border = thin_border

        c_fine = ws_skus.cell(row=row_idx, column=8, value=res.fines_ceiling_eur)
        c_fine.number_format = "#,##0 €"
        c_fine.border = thin_border

        ws_skus.cell(row=row_idx, column=9, value=res.summary).border = thin_border

    _auto_fit_columns(ws_summary, min_width=14, max_width=45)
    _auto_fit_columns(ws_skus, min_width=12, max_width=50)

    output_stream = io.BytesIO()
    wb.save(output_stream)
    return output_stream.getvalue()
